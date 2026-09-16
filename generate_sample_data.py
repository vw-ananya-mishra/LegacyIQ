#!/usr/bin/env python3
"""
Generate sample Legacy_Modernization_Synthetic_Dataset.xlsx for testing
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import random
from datetime import datetime, timedelta

def create_sample_xlsx(output_path="Legacy_Modernization_Synthetic_Dataset.xlsx"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    # Define styles
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    def add_header_row(ws, headers):
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Sheet 1: Applications
    ws_apps = wb.create_sheet("Applications")
    apps = ["OrderManagement", "BillingSystem", "InventoryTracker", "CustomerPortal", "AnalyticsEngine"]
    app_headers = ["Application_ID", "Name", "Description", "Status", "Criticality", "Complexity", 
                   "Technology_Stack", "Lines_of_Code", "Last_Modified"]
    add_header_row(ws_apps, app_headers)
    
    for idx, app in enumerate(apps, start=1):
        ws_apps.append([
            f"APP{idx:03d}",
            app,
            f"Legacy {app} system - mission critical",
            random.choice(["Active", "Maintenance", "Deprecated"]),
            random.choice(["critical", "high", "medium"]),
            random.choice(["high", "medium", "low"]),
            random.choice(["Java 6", "Python 2.7", "VB.NET", "C++", "COBOL"]),
            random.randint(10000, 500000),
            (datetime.now() - timedelta(days=random.randint(1, 365))).strftime("%Y-%m-%d")
        ])
    
    # Sheet 2: Code Modules
    ws_modules = wb.create_sheet("Code_Modules")
    module_headers = ["Module_ID", "Application_ID", "Name", "Description", "Language", 
                      "Lines_of_Code", "Complexity_Score", "Test_Coverage", "Critical_Functions", "Last_Modified"]
    add_header_row(ws_modules, module_headers)
    
    module_count = 0
    for app_idx in range(1, len(apps) + 1):
        for mod_idx in range(random.randint(3, 8)):
            module_count += 1
            ws_modules.append([
                f"MOD{module_count:04d}",
                f"APP{app_idx:03d}",
                f"Module_{mod_idx}",
                f"Legacy module with complex business logic",
                random.choice(["Java", "Python", "VB.NET", "C++"]),
                random.randint(1000, 50000),
                round(random.uniform(3, 9.5), 1),
                round(random.uniform(0, 100), 1),
                random.randint(0, 50),
                (datetime.now() - timedelta(days=random.randint(1, 365))).strftime("%Y-%m-%d")
            ])
    
    # Sheet 3: Business Rules
    ws_rules = wb.create_sheet("Business_Rules")
    rule_headers = ["Rule_ID", "Application_ID", "Rule_Name", "Description", "Rule_Type", 
                    "Condition", "Action", "Criticality", "Test_Coverage", "Last_Modified"]
    add_header_row(ws_rules, rule_headers)
    
    rule_count = 0
    rule_types = ["Validation", "Calculation", "Workflow", "Integration", "Compliance"]
    for app_idx in range(1, len(apps) + 1):
        for rule_idx in range(random.randint(5, 15)):
            rule_count += 1
            ws_rules.append([
                f"RULE{rule_count:04d}",
                f"APP{app_idx:03d}",
                f"Rule_{rule_idx}",
                f"Business rule for {random.choice(['validation', 'calculation', 'workflow'])}",
                random.choice(rule_types),
                f"IF condition THEN {random.choice(['approve', 'reject', 'escalate'])}",
                f"Execute action with consequences",
                random.choice(["critical", "high", "medium", "low"]),
                random.choice(["Full", "Partial", "None"]),
                (datetime.now() - timedelta(days=random.randint(1, 365))).strftime("%Y-%m-%d")
            ])
    
    # Sheet 4: Dependencies
    ws_deps = wb.create_sheet("Dependencies")
    dep_headers = ["Dependency_ID", "Source_ID", "Source_Type", "Target_ID", "Target_Type", 
                   "Relationship_Type", "Criticality", "Description"]
    add_header_row(ws_deps, dep_headers)
    
    dep_count = 0
    rel_types = ["calls", "depends_on", "uses_data_from", "integrates_with", "triggers"]
    source_types = ["module", "application", "data_store"]
    
    for _ in range(25):
        dep_count += 1
        ws_deps.append([
            f"DEP{dep_count:04d}",
            f"MOD{random.randint(1, module_count):04d}",
            "module",
            random.choice([f"APP{random.randint(1, len(apps)):03d}", f"DS{random.randint(1, 5):02d}"]),
            random.choice(["application", "data_store"]),
            random.choice(rel_types),
            random.choice(["critical", "high", "medium", "low"]),
            "Dependency relationship"
        ])
    
    # Sheet 5: Data Stores
    ws_datastores = wb.create_sheet("Data_Stores")
    ds_headers = ["DataStore_ID", "Name", "Store_Type", "Technology", "Criticality", 
                  "PII_Data", "Data_Volume", "Last_Modified"]
    add_header_row(ws_datastores, ds_headers)
    
    datastores = ["LegacyDB", "CacheLayer", "FileStore", "DataWarehouse", "MessageQueue"]
    for idx, ds in enumerate(datastores, start=1):
        ws_datastores.append([
            f"DS{idx:02d}",
            ds,
            random.choice(["Database", "File System", "Cache", "Message Queue"]),
            random.choice(["Oracle", "SQL Server", "DB2", "PostgreSQL", "Redis"]),
            random.choice(["critical", "high", "medium"]),
            random.choice([True, False]),
            random.choice(["10GB", "100GB", "1TB", "10TB"]),
            (datetime.now() - timedelta(days=random.randint(1, 365))).strftime("%Y-%m-%d")
        ])
    
    # Sheet 6: Integrations
    ws_integrations = wb.create_sheet("Integrations")
    int_headers = ["Integration_ID", "Application_ID", "External_System", "Integration_Type", 
                   "API_Version", "Criticality", "Status", "Last_Verified"]
    add_header_row(ws_integrations, int_headers)
    
    int_count = 0
    for app_idx in range(1, len(apps) + 1):
        for _ in range(random.randint(1, 4)):
            int_count += 1
            ws_integrations.append([
                f"INT{int_count:04d}",
                f"APP{app_idx:03d}",
                random.choice(["PaymentGateway", "EmailService", "Analytics", "Reporting", "ERP"]),
                random.choice(["REST", "SOAP", "FTP", "Database", "File"]),
                random.choice(["v1", "v2", "v3", "v4"]),
                random.choice(["critical", "high", "medium", "low"]),
                random.choice(["active", "deprecated", "retired"]),
                (datetime.now() - timedelta(days=random.randint(1, 365))).strftime("%Y-%m-%d")
            ])
    
    # Sheet 7: Test Cases
    ws_tests = wb.create_sheet("Test_Cases")
    test_headers = ["Test_ID", "Rule_ID", "Module_ID", "Test_Name", "Description", 
                    "Input_Data", "Expected_Output", "Test_Type", "Status", "Last_Run"]
    add_header_row(ws_tests, test_headers)
    
    test_count = 0
    for rule_idx in range(1, min(rule_count + 1, 25)):
        for _ in range(random.randint(0, 3)):
            test_count += 1
            ws_tests.append([
                f"TEST{test_count:04d}",
                f"RULE{rule_idx:04d}",
                f"MOD{random.randint(1, min(module_count, 30)):04d}",
                f"Test_{test_count}",
                f"Unit test for rule validation",
                "input_data",
                "expected_output",
                random.choice(["unit", "integration", "system"]),
                random.choice(["passed", "failed", "pending"]),
                (datetime.now() - timedelta(days=random.randint(1, 90))).strftime("%Y-%m-%d")
            ])
    
    # Sheet 8: Modernization Backlog
    ws_backlog = wb.create_sheet("Modernization_Backlog")
    backlog_headers = ["ID", "Application_ID", "Module_ID", "Recommendation", "Priority", 
                       "Effort_Estimate", "Risk_Level", "Dependencies", "Status"]
    add_header_row(ws_backlog, backlog_headers)
    
    backlog_count = 0
    for app_idx in range(1, len(apps) + 1):
        for _ in range(random.randint(2, 5)):
            backlog_count += 1
            ws_backlog.append([
                f"REC{backlog_count:04d}",
                f"APP{app_idx:03d}",
                f"MOD{random.randint(1, min(module_count, 30)):04d}",
                random.choice(["Refactor to microservices", "Migrate to cloud", "Modernize database", 
                              "Update UI framework", "Implement API gateway"]),
                random.choice(["high", "medium", "low"]),
                random.choice(["1-2 weeks", "2-4 weeks", "1-2 months", "2-3 months"]),
                random.choice(["high", "medium", "low"]),
                "APP001,DS01",
                random.choice(["draft", "pending_approval", "approved"])
            ])
    
    # Sheet 9: Documentation Artifacts
    ws_docs = wb.create_sheet("Documentation_Artifacts")
    doc_headers = ["ID", "Application_ID", "Artifact_Type", "Title", "Content_Summary", 
                   "Generated_Date", "Status", "Review_Status"]
    add_header_row(ws_docs, doc_headers)
    
    for app_idx in range(1, len(apps) + 1):
        for doc_type in ["Functional Spec", "Requirements", "API Spec"]:
            ws_docs.append([
                f"DOC{app_idx:03d}_{doc_type.replace(' ', '_')}",
                f"APP{app_idx:03d}",
                doc_type,
                f"{doc_type} for {apps[app_idx-1]}",
                "Auto-generated from source data",
                datetime.now().strftime("%Y-%m-%d"),
                "draft",
                "pending"
            ])
    
    # Sheet 10: README
    ws_readme = wb.create_sheet("README")
    ws_readme.append(["Field", "Description", "Example"])
    ws_readme.append(["Application_ID", "Unique identifier for application", "APP001"])
    ws_readme.append(["Criticality", "Business criticality level", "critical, high, medium, low"])
    ws_readme.append(["Complexity_Score", "Code complexity 0-10", "7.5"])
    ws_readme.append(["Test_Coverage", "Percentage of code covered by tests", "45.5"])
    ws_readme.append(["Status", "Current status of entity", "Active, Maintenance, Deprecated"])
    
    # Save the workbook
    wb.save(output_path)
    print(f"✓ Sample XLSX created: {output_path}")
    print(f"  - {len(apps)} applications")
    print(f"  - {module_count} modules")
    print(f"  - {rule_count} business rules")
    print(f"  - {dep_count} dependencies")
    print(f"  - {len(datastores)} data stores")
    print(f"  - {int_count} integrations")
    print(f"  - {test_count} test cases")
    print(f"  - {backlog_count} modernization items")

if __name__ == "__main__":
    create_sample_xlsx()
    print("\nYou can now upload this file to LegacyX!")
